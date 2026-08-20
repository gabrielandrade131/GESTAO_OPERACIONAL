from datetime import date

from django.test import SimpleTestCase
from docx import Document

from GO.proposal_official_pdf import (
    _apply_offshore_signature_details,
    resolve_official_template_path,
)


class OffshoreDocumentTests(SimpleTestCase):
    def test_replaces_signature_details_and_removes_the_stamp_image(self):
        document = Document(resolve_official_template_path("pc_offshore"))

        _apply_offshore_signature_details(
            document,
            "ANALISTA TESTE OFFSHORE",
            date(2026, 8, 19),
        )

        self.assertEqual(document.paragraphs[78].text, "Rio de Janeiro, 19 de agosto de 2026.")
        self.assertEqual(document.paragraphs[82].text, "ANALISTA TESTE OFFSHORE")
        self.assertNotIn("w:drawing", document.paragraphs[79]._p.xml)

        footer_text = "\n".join(
            paragraph.text
            for section in document.sections
            for paragraph in section.footer.paragraphs
        )
        self.assertIn("ANALISTA TESTE OFFSHORE - Comercial", footer_text)
        self.assertNotIn("Fernanda Braz", footer_text)
        analyst_footer = next(
            paragraph.text
            for section in document.sections
            for paragraph in section.footer.paragraphs
            if "ANALISTA TESTE OFFSHORE" in paragraph.text
        )
        self.assertTrue(analyst_footer.startswith("\t\t"))
