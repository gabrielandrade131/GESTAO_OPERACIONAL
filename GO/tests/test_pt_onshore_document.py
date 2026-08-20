from tempfile import TemporaryDirectory
from pathlib import Path

from django.test import SimpleTestCase
from docx import Document
from docx.oxml.ns import qn

from GO.proposal_official_pdf import _apply_pt_onshore_revision, resolve_official_template_path


class PtOnshoreDocumentTests(SimpleTestCase):
    serialized = {
        "numeroProposta": "PT-TESTE-123",
        "servico": "SERVIÇO TESTE PT 123",
        "empresa": "CLIENTE TESTE PT 123",
        "solicitante": "SOLICITANTE TESTE PT",
        "emailSolicitante": "pt.teste@example.com",
        "responsavel": "ANALISTA TESTE PT",
    }

    def build_document(self, conteudo):
        document = Document(resolve_official_template_path("pt_onshore"))
        _apply_pt_onshore_revision(
            document,
            {"conteudo": conteudo},
            self.serialized,
            show_variable_highlights=False,
        )
        with TemporaryDirectory() as directory:
            path = Path(directory) / "pt-onshore.docx"
            document.save(path)
            return Document(path)

    def document_text(self, document):
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    def test_replaces_only_mapped_pt_highlights(self):
        document = self.build_document({
            "resumo_planta": "RESUMO PLANTA TESTE PT 123",
            "referencias": ["REFERÊNCIA TESTE PT 123"],
            "sem_referencias": False,
            "prazo_execucao": "99 DIAS TESTE PT",
            "jornada": "08:00 ÀS 18:00 TESTE",
            "data_emissao": "2026-08-19",
        })
        text = self.document_text(document)

        for expected in (
            "PT-TESTE-123",
            "SERVIÇO TESTE PT 123",
            "CLIENTE TESTE PT 123",
            "RESUMO PLANTA TESTE PT 123",
            "REFERÊNCIA TESTE PT 123",
            "99 DIAS TESTE PT",
            "08:00 ÀS 18:00 TESTE",
        ):
            self.assertIn(expected, text)
        self.assertNotIn("xxxxx", text)
        self.assertNotIn("TIRAR SE NÃO HOUVER REFERÊNCIAS", text)

    def test_removes_reference_only_paragraphs_when_none_exist(self):
        document = self.build_document({
            "resumo_planta": "RESUMO PT",
            "referencias": [],
            "sem_referencias": True,
            "prazo_execucao": "30 dias",
            "jornada": "07:00 às 17:00.",
            "data_emissao": "2026-08-19",
        })
        text = self.document_text(document)

        self.assertNotIn("xxxxx", text)
        self.assertNotIn("TIRAR SE NÃO HOUVER REFERÊNCIAS", text)
        self.assertNotIn("Para a elaboração da proposta Técnica utilizamos como referência", text)

    def test_renders_each_pt_reference_in_a_separate_styled_paragraph(self):
        references = ["REFERÊNCIA A", "REFERÊNCIA B", "REFERÊNCIA C"]
        document = self.build_document({
            "resumo_planta": "RESUMO PT",
            "referencias": references,
            "sem_referencias": False,
            "prazo_execucao": "30 dias",
            "jornada": "07:00 as 17:00.",
            "data_emissao": "2026-08-19",
        })
        rendered = [paragraph for paragraph in document.paragraphs if paragraph.text in references]

        self.assertEqual([paragraph.text for paragraph in rendered], references)
        self.assertTrue(all("\n" not in paragraph.text for paragraph in rendered))
        self.assertEqual(rendered[0]._p.pPr.xml, rendered[1]._p.pPr.xml)

    def test_grows_the_official_revision_table_for_revision_history(self):
        document = self.build_document({
            "resumo_planta": "RESUMO PT",
            "referencias": ["REFERÊNCIA"],
            "sem_referencias": False,
            "prazo_execucao": "30 dias",
            "jornada": "07:00 às 17:00.",
            "data_emissao": "2026-08-19",
            "quadro_revisoes": [
                {"revisao": "00", "data": "2026-08-18", "descricao": "Emissão Inicial."},
                {"revisao": "01", "data": "2026-08-19", "descricao": "Ajuste de escopo."},
            ],
        })
        table = document.tables[1]

        self.assertEqual(len(table.rows), 3)
        self.assertEqual(table.rows[1].cells[0].text, "00")
        self.assertEqual(table.rows[2].cells[0].text, "01")
        self.assertEqual(table.rows[2].cells[2].text, "Ajuste de escopo.")

    def test_fills_labor_histogram_with_dynamic_template_rows(self):
        document = self.build_document({
            "resumo_planta": "RESUMO PT",
            "referencias": ["REFERÊNCIA"],
            "sem_referencias": False,
            "prazo_execucao": "30 dias",
            "jornada": "07:00 às 17:00.",
            "data_emissao": "2026-08-19",
            "histograma_mao_obra": [
                {"funcao": "Supervisor", "quantidade": "1"},
                {"funcao": "Ajudante operacional", "quantidade": "2"},
            ],
        })
        table = document.tables[2]

        self.assertEqual(len(table.rows), 3)
        self.assertEqual(table.rows[1].cells[0].text, "Supervisor")
        self.assertEqual(table.rows[1].cells[1].text, "1")
        self.assertEqual(table.rows[2].cells[0].text, "Ajudante operacional")
        self.assertEqual(table.rows[2].cells[1].text, "2")

    def test_inserts_optional_methodology_before_the_histogram(self):
        document = self.build_document({
            "resumo_planta": "RESUMO PT",
            "metodologia_executiva": "METODOLOGIA TESTE PT 123",
            "referencias": ["REFERENCIA"],
            "sem_referencias": False,
            "prazo_execucao": "30 dias",
            "jornada": "07:00 as 17:00.",
            "data_emissao": "2026-08-19",
        })
        paragraphs = [paragraph.text for paragraph in document.paragraphs]

        self.assertLess(paragraphs.index("METODOLOGIA TESTE PT 123"), paragraphs.index("8.2 HISTOGRAMA"))

    def test_inserts_optional_histogram_description_before_labor_histogram(self):
        document = self.build_document({
            "resumo_planta": "RESUMO PT",
            "referencias": ["REFERENCIA"],
            "sem_referencias": False,
            "prazo_execucao": "30 dias",
            "jornada": "07:00 as 17:00.",
            "data_emissao": "2026-08-19",
            "descricao_histograma": "DESCRIÇÃO DO HISTOGRAMA TESTE PT 123",
        })
        paragraphs = [paragraph.text for paragraph in document.paragraphs]

        self.assertLess(paragraphs.index("8.2 HISTOGRAMA"), paragraphs.index("DESCRIÇÃO DO HISTOGRAMA TESTE PT 123"))
        self.assertLess(paragraphs.index("DESCRIÇÃO DO HISTOGRAMA TESTE PT 123"), paragraphs.index("8.2.1 HISTOGRAMA DE MÃO DE OBRA"))

    def test_fills_equipment_histogram_with_dynamic_template_rows(self):
        document = self.build_document({
            "resumo_planta": "RESUMO PT",
            "referencias": ["REFERENCIA"],
            "sem_referencias": False,
            "prazo_execucao": "30 dias",
            "jornada": "07:00 as 17:00.",
            "data_emissao": "2026-08-19",
            "histograma_equipamentos": ["Bomba pneumática", "Detector de gases"],
        })

        table = document.tables[3]
        self.assertEqual(len(table.rows), 3)
        self.assertEqual(table.rows[1].cells[0].text, "Bomba pneumática")
        self.assertEqual(table.rows[2].cells[0].text, "Detector de gases")

    def test_replaces_pt_footer_date_and_analyst_without_signature_image(self):
        document = self.build_document({
            "resumo_planta": "RESUMO PT",
            "referencias": ["REFERENCIA"],
            "sem_referencias": False,
            "prazo_execucao": "30 dias",
            "jornada": "07:00 as 17:00.",
            "data_emissao": "2026-08-19",
        })
        text = self.document_text(document)

        self.assertIn("Agosto / 2026", text)
        self.assertIn("Rio de Janeiro, 19 de agosto de 2026.", text)
        self.assertIn("ANALISTA TESTE PT", text)
        self.assertNotIn("Katlyn Brito", text)
        analyst_index = next(
            index for index, paragraph in enumerate(document.paragraphs)
            if paragraph.text == "ANALISTA TESTE PT"
        )
        self.assertFalse("w:drawing" in document.paragraphs[analyst_index - 1]._p.xml)
        footer_text = "\n".join(
            paragraph.text
            for section in document.sections
            for paragraph in section.footer.paragraphs
        )
        self.assertIn("ANALISTA TESTE PT – Comercial", footer_text)
        self.assertNotIn("Katlyn Brito", footer_text)

    def test_aligns_pt_footer_logos_at_the_left_margin(self):
        document = self.build_document({
            "resumo_planta": "RESUMO PT",
            "referencias": ["REFERENCIA"],
            "sem_referencias": False,
            "prazo_execucao": "30 dias",
            "jornada": "07:00 as 17:00.",
            "data_emissao": "2026-08-19",
        })

        positions = []
        for section in document.sections:
            for paragraph in section.footer.paragraphs:
                for anchor in paragraph._p.xpath(".//wp:anchor"):
                    if anchor.find(".//" + qn("a:blip")) is None:
                        continue
                    position = anchor.find(qn("wp:positionH"))
                    positions.append((position.get("relativeFrom"), position.find(qn("wp:posOffset")).text))

        self.assertTrue(positions)
        self.assertEqual(positions, [("margin", "0")] * len(positions))
