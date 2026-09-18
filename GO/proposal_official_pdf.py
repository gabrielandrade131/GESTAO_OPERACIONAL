"""Generate official Commercial proposal PDFs from the approved DOCX templates."""

from __future__ import annotations

from copy import deepcopy
from datetime import date
from decimal import Decimal
from pathlib import Path
import os
import re
import shutil
import subprocess
import tempfile

from django.conf import settings
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


class OfficialProposalPdfError(Exception):
    """Raised when an official proposal cannot be generated safely."""


TEMPLATE_DIR = Path(settings.BASE_DIR) / "GO" / "docs" / "propostas_oficiais"
TEMPLATE_FILES = {
    "pc_offshore": Path("Offshore") / "xxxx.xx.xx - Proposta padrão - Offshore.docx",
    "pc_onshore": Path("Onshore_PC") / "XXXX.XX.XX_PC - Onshore.docx",
    "pt_onshore": Path("Onshore_PT") / "XXXX.XX.XX_PT - Onshore.docx",
}

PORTUGUESE_MONTHS = (
    "janeiro", "fevereiro", "março", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
)


# Prefer the organized document directories, while keeping a safe fallback for
# deployments that have not yet moved the legacy templates from the root.
TEMPLATE_ROOT_FALLBACKS = {
    "pc_offshore": "xxxx.xx.xx - Proposta padrão - Offshore.docx",
    "pc_onshore": "XXXX.XX.XX_PC - Onshore.docx",
    "pt_onshore": "XXXX.XX.XX_PT - Onshore.docx",
}


def _clean(value):
    return str(value or "").strip()


def _revision_number(value):
    try:
        return int(str(value or "0"))
    except (TypeError, ValueError):
        return 0


def resolve_official_template_path(template_key):
    """Resolve a bundled official template without using an absolute path."""
    try:
        preferred_path = TEMPLATE_DIR / TEMPLATE_FILES[template_key]
        legacy_name = TEMPLATE_ROOT_FALLBACKS[template_key]
    except KeyError as exc:
        raise OfficialProposalPdfError("O tipo de documento oficial informado não é suportado.") from exc

    if preferred_path.exists():
        return preferred_path

    legacy_path = TEMPLATE_DIR / legacy_name
    if legacy_path.exists():
        return legacy_path

    raise OfficialProposalPdfError(
        f"O template oficial '{legacy_name}' não foi encontrado em {TEMPLATE_DIR}."
    )


def _format_currency(value):
    amount = Decimal(str(value or "0"))
    return "R$ {:,.2f}".format(amount).replace(",", "X").replace(".", ",").replace("X", ".")


def _format_date_long(value):
    if not isinstance(value, date):
        return ""
    return f"Rio de Janeiro, {value.day:02d} de {PORTUGUESE_MONTHS[value.month - 1]} de {value.year}."


def _format_cover_month(value):
    if not isinstance(value, date):
        return ""
    return f"{PORTUGUESE_MONTHS[value.month - 1].capitalize()} / {value.year}"


def _format_document_reference_date(value):
    """Format the manually reviewed technical-proposal date for the PC text."""
    raw_value = _clean(value)
    if not raw_value:
        return ""
    try:
        parsed = date.fromisoformat(raw_value)
    except ValueError:
        return raw_value
    return f"{parsed.day} de {PORTUGUESE_MONTHS[parsed.month - 1]} de {parsed.year}"


def _proposal_kind(proposal):
    operation = _clean(getattr(getattr(proposal, "tipo_operacao", None), "tipo_operacao", "")).casefold()
    is_offshore = operation == "offshore"
    is_onshore = operation == "onshore"
    if not (is_offshore or is_onshore):
        raise OfficialProposalPdfError("Informe o tipo de operação Onshore ou Offshore para gerar a proposta oficial.")

    # PC has precedence because existing records can carry both technical and
    # commercial control values while still representing a commercial proposal.
    # Older commercial records persist the label in the feminine form
    # ("Elaborada"), while newer records use the field's canonical masculine
    # value ("Elaborado"). Both mean that the commercial document is ready.
    if _clean(proposal.pc_ptc).casefold() in {"elaborado", "elaborada"}:
        return "pc_offshore" if is_offshore else "pc_onshore"
    if _clean(proposal.pt_financeiro).casefold() == "elaborada":
        if is_offshore:
            raise OfficialProposalPdfError("Ainda não existe template oficial de Proposta Técnica Offshore.")
        return "pt_onshore"

    raise OfficialProposalPdfError(
        "Defina PC / PTC como Elaborado ou PT como Elaborada para identificar o template oficial."
    )


def _paragraphs(container):
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _paragraphs(cell)


def _replace_in_runs(paragraph, source, target):
    """Replace split DOCX runs while retaining the style of the first run."""
    if not source or source == target or source not in paragraph.text:
        return

    while source in paragraph.text:
        runs = paragraph.runs
        combined = "".join(run.text for run in runs)
        start = combined.find(source)
        end = start + len(source)
        cursor = 0
        start_run = end_run = None
        start_offset = end_offset = 0
        for index, run in enumerate(runs):
            next_cursor = cursor + len(run.text)
            if start_run is None and start < next_cursor:
                start_run, start_offset = index, start - cursor
            if end <= next_cursor:
                end_run, end_offset = index, end - cursor
                break
            cursor = next_cursor
        if start_run is None or end_run is None:
            return

        prefix = runs[start_run].text[:start_offset]
        suffix = runs[end_run].text[end_offset:]
        if start_run == end_run:
            runs[start_run].text = f"{prefix}{target}{suffix}"
        else:
            runs[start_run].text = f"{prefix}{target}"
            for index in range(start_run + 1, end_run):
                runs[index].text = ""
            runs[end_run].text = suffix


def _replace_document_text(document, replacements):
    containers = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))
    for container in containers:
        for paragraph in _paragraphs(container):
            for source, target in replacements.items():
                _replace_in_runs(paragraph, source, target)


def _remove_text_highlights(document):
    """Remove Word text highlights from the issued PDF without touching table fills."""
    containers = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))
    for container in containers:
        for paragraph in _paragraphs(container):
            for run in paragraph.runs:
                properties = run._r.rPr
                if properties is None:
                    continue
                for highlight in list(properties.findall(qn("w:highlight"))):
                    properties.remove(highlight)


def _clear_run_highlight(run):
    """Remove only the editing marker from a mapped variable DOCX run."""
    properties = run._r.rPr
    if properties is None:
        return
    for highlight in list(properties.findall(qn("w:highlight"))):
        properties.remove(highlight)


def _set_highlighted_runs(paragraph, value, *, final):
    """Replace the yellow runs in a known paragraph without touching its other text."""
    marked_runs = [run for run in paragraph.runs if run.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    if not marked_runs:
        raise OfficialProposalPdfError("O campo variavel esperado nao foi encontrado no template PT Onshore.")
    marked_runs[0].text = _clean(value)
    for run in marked_runs[1:]:
        run.text = ""
    if final:
        for run in marked_runs:
            _clear_run_highlight(run)


def _remove_paragraph(paragraph):
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)


def _format_revision_table_date(value):
    if not isinstance(value, date):
        return ""
    abbreviations = ("jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez")
    return f"{value.day:02d}-{abbreviations[value.month - 1]}-{value.year}"


def _parse_review_date(value):
    raw_value = _clean(value)
    if not raw_value:
        return None
    try:
        return date.fromisoformat(raw_value)
    except ValueError as error:
        raise OfficialProposalPdfError("Informe uma data valida para a Proposta Tecnica.") from error


def _set_cell_text(cell, value):
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = _clean(value)
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(_clean(value))


def _set_row_cells(row, values):
    """Fill a row without replacing its XML, preserving the official table style."""
    for column, value in enumerate(values):
        if column < len(row.cells):
            _set_cell_text(row.cells[column], value)


def _add_cell_logo(cell):
    """Place the approved Ambipar mark in the PPU header when the asset exists."""
    logo_path = Path(settings.BASE_DIR) / "GO" / "static" / "js" / "img" / "Logo_Preto.png"
    if not logo_path.exists():
        return
    paragraph = cell.paragraphs[0]
    for run in paragraph.runs:
        run.text = ""
    paragraph.alignment = 1
    paragraph.add_run().add_picture(str(logo_path), width=Inches(0.85))


def _set_cell_shading(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_ppu_cell_text(cell, value, *, size, bold=False, color="000000", alignment=None):
    """Set PPU content with the typography used by the official price sheet."""
    _set_cell_text(cell, value)
    paragraph = cell.paragraphs[0]
    if alignment is not None:
        paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in paragraph.runs:
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)


def _style_ppu_row(row, *, size, bold=False, color="000000", alignments=()):
    seen = set()
    for index, cell in enumerate(row.cells):
        if id(cell._tc) in seen:
            continue
        seen.add(id(cell._tc))
        _set_ppu_cell_text(
            cell,
            cell.text,
            size=size,
            bold=bold,
            color=color,
            alignment=alignments[index] if index < len(alignments) else WD_ALIGN_PARAGRAPH.CENTER,
        )


def _fill_pc_onshore_financial_table(document, items, revision, serialized):
    """Populate the blank PPU table in the official PC Onshore template."""
    if len(document.tables) < 2:
        raise OfficialProposalPdfError("A tabela financeira do template PC Onshore não foi encontrada.")

    table = document.tables[1]
    if len(table.rows) < 5:
        raise OfficialProposalPdfError("A estrutura da tabela PPU do template PC Onshore está incompleta.")

    content = (revision or {}).get("conteudo") or {}
    selected = content.get("itens_financeiros") or {}
    selected_items = []
    for item in items:
        selection = selected.get(str(item.get("id"))) or {}
        if selection.get("incluir") is False:
            continue
        quantity = Decimal(str(selection.get("quantidade") or item.get("quantidade") or "1"))
        if quantity <= 0:
            continue
        selected_items.append({**item, "quantidade": quantity, "subtotal": item["preco_unitario"] * quantity})

    if not selected_items:
        raise OfficialProposalPdfError("Selecione ao menos um item financeiro para preencher a PPU.")

    total = sum((item["subtotal"] for item in selected_items), Decimal("0"))
    scope = _clean(content.get("escopo_ppu")) or _clean(serialized.get("servico") or serialized.get("escopo"))
    deadline = _clean(content.get("prazo_ppu")) or _clean(content.get("prazo")) or _clean(serialized.get("tempoContratoDias"))
    number = _clean(serialized.get("numeroProposta"))

    _set_cell_shading(table.rows[0].cells[0], "D4FF00")
    _set_ppu_cell_text(table.rows[0].cells[0], "PLANILHA DE PREÇOS UNITÁRIOS - PPU", size=9, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    # The first PPU rows contain merged cells. Write each physical cell only
    # once so a later logical column cannot erase the preceding label.
    _set_ppu_cell_text(table.rows[1].cells[0], "Cliente:", size=8, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _set_ppu_cell_text(table.rows[1].cells[1], _clean(serialized.get("empresa")), size=8, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _set_ppu_cell_text(table.rows[1].cells[2], f"ID: {number}", size=8, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    _add_cell_logo(table.rows[1].cells[-1])
    _set_ppu_cell_text(table.rows[2].cells[0], "ESCOPO:", size=8, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _set_ppu_cell_text(table.rows[2].cells[1], scope, size=8, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _set_ppu_cell_text(table.rows[2].cells[2], f"Prazo: {deadline}" if deadline else "Prazo:", size=8, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    _set_ppu_cell_text(table.rows[2].cells[4], "Preço Total:", size=8, bold=False, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    _set_ppu_cell_text(table.rows[2].cells[5], _format_currency(total), size=9, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    _set_row_cells(table.rows[3], ("ITENS", "EXECUÇÃO DOS SERVIÇOS", "UNID.", "QUANT.", "CUSTO UNITÁRIO", "VALOR TOTAL"))
    _style_ppu_row(table.rows[3], size=7, bold=True, color="FFFFFF")

    reference_row = deepcopy(table.rows[4]._tr)
    for row in list(table.rows[4:]):
        table._tbl.remove(row._tr)

    for index, item in enumerate(selected_items, start=1):
        table._tbl.append(deepcopy(reference_row))
        row = table.rows[-1]
        _set_row_cells(
            row,
            (
                str(index),
                item["label"],
                _clean(item.get("unidade")) or "UN",
                str(item["quantidade"]).replace(".", ","),
                _format_currency(item["preco_unitario"]),
                _format_currency(item["subtotal"]),
            ),
        )
        _style_ppu_row(
            row,
            size=8,
            alignments=(
                WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.CENTER,
            ),
        )


def _apply_pc_onshore_revision(document, revision, serialized, items, *, show_variable_highlights=False):
    """Apply only the PC Onshore review content to its official DOCX template."""
    content = (revision or {}).get("conteudo") or {}
    paragraphs = document.paragraphs
    if len(paragraphs) > 55:
        service = _clean(serialized.get("servico") or serialized.get("escopo"))
        if service:
            presentation = (
                "É com satisfação que a AMBIPAR RESPONSE TANK CLEANING apresenta sua "
                "Proposta Comercial para prestação de serviços especializados de "
                f"{service.casefold()}."
            )
            service_complement = _clean(content.get("complemento_servico"))
            if service_complement:
                presentation = f"{presentation} {service_complement}"
            _set_paragraph_text(paragraphs[22], presentation)

        reference_date = _format_document_reference_date(content.get("pt_data"))
        if not reference_date:
            raise OfficialProposalPdfError("Informe a data da Proposta Técnica antes de gerar o documento.")
        reference = " - ".join(
            value for value in (
                f"{_clean(serialized.get('numeroProposta'))} PT",
                _clean(serialized.get("empresa")),
                _clean(serialized.get("servico") or serialized.get("escopo")),
                _clean(serialized.get("unidade")),
            ) if value
        )
        _set_paragraph_with_highlight(
            paragraphs[37],
            "Para esta proposta comercial o escopo da proposta, detalhamentos técnicos, prazo de "
            "execução e regime de trabalho, premissas técnicas, obrigações da contratante, obrigações "
            "da contratada e itens fora do escopo estão alinhados com a Proposta Técnica “",
            reference,
            f"” datada de {reference_date}, no qual um documento complementa o outro e não devem ser analisados separadamente.",
            show_highlight=show_variable_highlights,
        )

        deadline = _clean(content.get("prazo"))
        if deadline:
            _set_paragraph_text(paragraphs[47], f"O prazo de execução será de {deadline}.")
        if _clean(content.get("prazo_complementar")):
            _set_paragraph_text(paragraphs[48], _clean(content.get("prazo_complementar")))
        validity = _clean(content.get("validade_dias"))
        if validity:
            _set_paragraph_text(paragraphs[55], f"As condições desta proposta são válidas por {validity} dias a partir da data de sua emissão.")

    _fill_pc_onshore_financial_table(document, items, revision, serialized)


def _append_pc_onshore_signature(document, analyst):
    """Clone the approved Offshore signature block for the PC Onshore layout."""
    date_paragraph = next(
        (
            paragraph
            for paragraph in reversed(document.paragraphs)
            if paragraph.text.strip().startswith("Rio de Janeiro,")
        ),
        None,
    )
    if date_paragraph is None:
        raise OfficialProposalPdfError("A data de assinatura do template PC Onshore nao foi encontrada.")

    offshore_template = Document(resolve_official_template_path("pc_offshore"))
    # Paragraphs 80-85 are the approved plain-text signature block. Paragraph
    # 79 contains the legacy signature drawing and must never be copied.
    source_block = offshore_template.paragraphs[80:86]
    if len(source_block) != 6:
        raise OfficialProposalPdfError("O bloco de assinatura do template Offshore nao foi encontrado.")

    insertion_point = date_paragraph
    inserted = []
    for source_paragraph in source_block:
        paragraph_xml = deepcopy(source_paragraph._p)
        insertion_point._p.addnext(paragraph_xml)
        insertion_point = Paragraph(paragraph_xml, date_paragraph._parent)
        inserted.append(insertion_point)

    # The third cloned paragraph is the analyst name. Reusing the original
    # paragraph preserves its exact approved formatting and alignment.
    _set_paragraph_text(inserted[2], (analyst or "Nao informado").upper())


def _set_pc_onshore_footer_analyst(document, analyst):
    """Keep the official PC footer tied to the commercial analyst."""
    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            if "Comercial" not in paragraph.text:
                continue
            _replace_in_runs(
                paragraph,
                "Nome do analista \u2013 Comercial",
                f"{analyst or 'Nao informado'} \u2013 Comercial",
            )


def _apply_pt_onshore_revision(document, revision, serialized, *, show_variable_highlights=False):
    """Apply the PT revision only to the explicitly yellow runs of its official DOCX."""
    content = (revision or {}).get("conteudo") or {}
    paragraphs = document.paragraphs
    if len(paragraphs) <= 148:
        raise OfficialProposalPdfError("A estrutura esperada do template PT Onshore nao foi encontrada.")

    service = _clean(serialized.get("servico") or serialized.get("escopo"))
    client = _clean(serialized.get("empresa"))
    requester = _clean(serialized.get("solicitante"))
    email = _clean(serialized.get("emailSolicitante"))
    number = _clean(serialized.get("numeroProposta"))
    analyst = _clean(serialized.get("responsavel"))
    summary = _clean(content.get("resumo_planta"))
    methodology = _clean(content.get("metodologia_executiva"))
    histogram_description = _clean(content.get("descricao_histograma"))
    deadline = _clean(content.get("prazo_execucao"))
    journey = _clean(content.get("jornada"))
    references = [
        _clean(reference)
        for reference in (content.get("referencias") or [])
        if _clean(reference)
    ]
    without_references = bool(content.get("sem_referencias"))
    document_date = _parse_review_date(content.get("data_emissao"))
    missing = [
        label for label, value in (
            ("resumo sobre a planta", summary),
            ("prazo de execucao", deadline),
            ("jornada", journey),
            ("data de emissao", document_date),
        ) if not value
    ]
    if missing:
        raise OfficialProposalPdfError("Preencha os campos da revisao PT: " + ", ".join(missing) + ".")
    if without_references and references:
        raise OfficialProposalPdfError("Escolha referencias ou marque que nao existem referencias especificas.")
    if not without_references and not references:
        raise OfficialProposalPdfError("Informe ao menos uma referencia ou marque que nao existem referencias especificas.")

    final = not show_variable_highlights
    histogram_note_template = deepcopy(paragraphs[89]._p)
    _set_highlighted_runs(paragraphs[1], number, final=final)
    _set_highlighted_runs(paragraphs[5], service, final=final)
    _set_highlighted_runs(paragraphs[8], client, final=final)
    _set_highlighted_runs(paragraphs[11], requester, final=final)
    _set_highlighted_runs(paragraphs[12], email, final=final)
    _set_highlighted_runs(paragraphs[40], f"A {client},", final=final)
    _set_highlighted_runs(paragraphs[41], f"de {service.casefold()}", final=final)
    _set_highlighted_runs(paragraphs[65], summary, final=final)
    _set_highlighted_runs(paragraphs[85], f"de {deadline}", final=final)
    if methodology:
        _set_paragraph_with_highlight(
            paragraphs[89],
            "",
            methodology,
            "",
            show_highlight=show_variable_highlights,
        )
    if histogram_description:
        histogram_note_xml = deepcopy(histogram_note_template)
        paragraphs[90]._p.addnext(histogram_note_xml)
        histogram_note = Paragraph(histogram_note_xml, paragraphs[90]._parent)
        _set_paragraph_with_highlight(
            histogram_note,
            "",
            histogram_description,
            "",
            show_highlight=show_variable_highlights,
        )
    _set_highlighted_runs(paragraphs[114], journey, final=final)
    _set_paragraph_with_highlight(
        paragraphs[15],
        "",
        _format_cover_month(document_date),
        "",
        show_highlight=show_variable_highlights,
    )
    _set_paragraph_with_highlight(
        paragraphs[148],
        "Rio de Janeiro, ",
        f"{document_date.day:02d} de {PORTUGUESE_MONTHS[document_date.month - 1]} de {document_date.year}",
        ".",
        show_highlight=show_variable_highlights,
    )
    _set_paragraph_text(paragraphs[151], analyst or "Não informado")
    _set_pt_onshore_footer_analyst(document, analyst)

    if without_references:
        _remove_paragraph(paragraphs[83])
        _remove_paragraph(paragraphs[105])
    else:
        _fill_pt_onshore_references(paragraphs[83], references, final=final)
        _set_highlighted_runs(
            paragraphs[105],
            "Para a elaboracao da proposta Tecnica utilizamos como referencia todos os documentos citados no capitulo REFERENCIAS, questionamentos enviados e experiencia da Ambipar em servicos similares prestados anteriormente;",
            final=final,
        )

    for section in document.sections:
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if any(run.font.highlight_color == WD_COLOR_INDEX.YELLOW for run in paragraph.runs):
                            _set_highlighted_runs(paragraph, service, final=final)

    revision_rows = content.get("quadro_revisoes") or [{
            "revisao": _clean(serialized.get("rev")) or "00",
            "data": content.get("data_revisao") or content.get("data_emissao"),
            "descricao": content.get("descricao_revisao") or "Emissão Inicial.",
        }]
    current_revision = max((_revision_number(row.get("revisao")) for row in revision_rows), default=0)
    current_row = next((row for row in revision_rows if _revision_number(row.get("revisao")) == current_revision), {})
    if current_revision > 0 and (not _clean(current_row.get("data")) or not _clean(current_row.get("descricao"))):
        raise OfficialProposalPdfError("Informe a data e a descricao da revisao atual antes de gerar o documento.")
    _fill_pt_onshore_revision_table(document, revision_rows, final=final)

    histogram_rows = []
    for row in content.get("histograma_mao_obra") or []:
        function = _clean(row.get("funcao"))
        quantity = _revision_number(row.get("quantidade"))
        if not function and not _clean(row.get("quantidade")):
            continue
        if not function or quantity <= 0 or str(quantity) != _clean(row.get("quantidade")):
            raise OfficialProposalPdfError("Informe função e quantidade inteira positiva para cada linha do histograma.")
        histogram_rows.append({"funcao": function, "quantidade": quantity})
    _fill_pt_onshore_histogram_table(document, histogram_rows)

    equipment_histogram_rows = [
        _clean(description)
        for description in content.get("histograma_equipamentos") or []
        if _clean(description)
    ]
    _fill_pt_onshore_equipment_histogram_table(document, equipment_histogram_rows)

    # The signature image belongs to the editable template only. The issued
    # proposal identifies the commercial analyst by name, without a signature.
    _remove_paragraph(paragraphs[149])
    _normalize_pt_onshore_footer_logo_positions(document)


def _fill_pt_onshore_revision_table(document, revision_rows, *, final):
    """Clone the official review-table row for every commercial revision."""
    if len(document.tables) <= 1 or len(document.tables[1].rows) <= 1:
        raise OfficialProposalPdfError("O quadro de revisoes do template PT Onshore nao foi encontrado.")
    table = document.tables[1]
    reference_row = deepcopy(table.rows[1]._tr)
    for row in list(table.rows[1:]):
        table._tbl.remove(row._tr)

    for source_row in revision_rows:
        raw_date = _clean(source_row.get("data"))
        try:
            formatted_date = _format_revision_table_date(date.fromisoformat(raw_date)) if raw_date else ""
        except ValueError:
            formatted_date = raw_date
        table._tbl.append(deepcopy(reference_row))
        row = table.rows[-1]
        _set_row_cells(
            row,
            (
                _clean(source_row.get("revisao")),
                formatted_date,
                _clean(source_row.get("descricao")),
            ),
        )
        if final:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _clear_run_highlight(run)


def _fill_pt_onshore_histogram_table(document, rows):
    """Fill the official labor histogram by cloning its blank template row."""
    if len(document.tables) <= 2 or len(document.tables[2].rows) <= 1:
        raise OfficialProposalPdfError("A tabela de histograma do template PT Onshore nao foi encontrada.")
    if not rows:
        return
    table = document.tables[2]
    reference_row = deepcopy(table.rows[1]._tr)
    for row in list(table.rows[1:]):
        table._tbl.remove(row._tr)
    for item in rows:
        table._tbl.append(deepcopy(reference_row))
        row = table.rows[-1]
        _set_row_cells(row, (_clean(item.get("funcao")), str(item.get("quantidade"))))


def _fill_pt_onshore_equipment_histogram_table(document, rows):
    """Fill the official equipment histogram by cloning its blank template row."""
    if len(document.tables) <= 3 or len(document.tables[3].rows) <= 1:
        raise OfficialProposalPdfError("A tabela de histograma de equipamentos do template PT Onshore nao foi encontrada.")
    if not rows:
        return
    table = document.tables[3]
    reference_row = deepcopy(table.rows[1]._tr)
    for row in list(table.rows[1:]):
        table._tbl.remove(row._tr)
    for description in rows:
        table._tbl.append(deepcopy(reference_row))
        _set_cell_text(table.rows[-1].cells[0], description)


def _set_pt_onshore_footer_analyst(document, analyst):
    """Replace only the template analyst name in each PT section footer."""
    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            if "Katlyn Brito" in paragraph.text:
                _replace_in_runs(paragraph, "Katlyn Brito", analyst or "Não informado")


def _fill_pt_onshore_references(paragraph, references, *, final):
    """Create one styled official-template paragraph for each PT reference."""
    template_xml = deepcopy(paragraph._p)
    current_paragraph = paragraph
    for index, reference in enumerate(references):
        if index:
            new_xml = deepcopy(template_xml)
            current_paragraph._p.addnext(new_xml)
            current_paragraph = Paragraph(new_xml, paragraph._parent)
        _set_highlighted_runs(current_paragraph, reference, final=final)


def _normalize_pt_onshore_footer_logo_positions(document):
    """Anchor every PT footer logo at the same left margin across sections."""
    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            for anchor in paragraph._p.xpath(".//wp:anchor"):
                if anchor.find(".//" + qn("a:blip")) is None:
                    continue
                position = anchor.find(qn("wp:positionH"))
                if position is None:
                    continue
                position.set("relativeFrom", "margin")
                offset = position.find(qn("wp:posOffset"))
                if offset is not None:
                    offset.text = "0"


def _fill_offshore_financial_table(document, items):
    if len(document.tables) < 6:
        raise OfficialProposalPdfError("A tabela financeira do template Offshore não foi encontrada.")

    table = document.tables[5]
    reference_row = deepcopy(table.rows[1]._tr) if len(table.rows) > 1 else None
    for row in list(table.rows[1:]):
        table._tbl.remove(row._tr)

    if not items:
        raise OfficialProposalPdfError("A proposta Offshore precisa possuir ao menos um item financeiro para gerar o PDF oficial.")

    for index, item in enumerate(items, start=1):
        if reference_row is not None:
            table._tbl.append(deepcopy(reference_row))
            row = table.rows[-1]
        else:
            row = table.add_row()
        _set_cell_text(row.cells[0], str(index))
        _set_cell_text(row.cells[1], item["label"])
        _set_cell_text(row.cells[2], "")
        _set_cell_text(row.cells[3], _format_currency(item["preco_unitario"]) if item["preco_unitario"] else "")


def _fill_table_rows(table, rows, columns):
    """Replace variable table rows while preserving the template row formatting."""
    reference_row = deepcopy(table.rows[1]._tr) if len(table.rows) > 1 else None
    for row in list(table.rows[1:]):
        table._tbl.remove(row._tr)
    for index, values in enumerate(rows, start=1):
        if reference_row is not None:
            table._tbl.append(deepcopy(reference_row))
            row = table.rows[-1]
        else:
            row = table.add_row()
        for column, value in enumerate(values[:columns]):
            _set_cell_text(row.cells[column], value)


def _set_paragraph_text(paragraph, value):
    if paragraph.runs:
        paragraph.runs[0].text = _clean(value)
        for run in paragraph.runs[1:]:
            run.text = ""
        return

    # Word's static table of contents stores its visible labels inside
    # hyperlinks, which python-docx does not expose through paragraph.runs.
    # Update only the first visible label and preserve the page-reference field.
    hyperlink_texts = paragraph._p.xpath(".//w:hyperlink//w:t")
    if hyperlink_texts:
        hyperlink_texts[0].text = _clean(value)
        return

    paragraph.add_run(_clean(value))


def _set_paragraph_with_highlight(paragraph, prefix, highlighted, suffix, *, show_highlight):
    """Replace a paragraph while preserving its base run formatting."""
    base_run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    base_properties = deepcopy(base_run._r.rPr) if base_run._r.rPr is not None else None
    # Prefixes/suffixes can deliberately contain spaces around a highlighted
    # variable (for example, a date after a comma), so do not normalize them.
    base_run.text = str(prefix or "")
    for run in paragraph.runs[1:]:
        run.text = ""

    variable_run = paragraph.add_run(_clean(highlighted))
    if base_properties is not None:
        variable_run._r.insert(0, deepcopy(base_properties))
    if show_highlight:
        variable_run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    ending_run = paragraph.add_run(str(suffix or ""))
    if base_properties is not None:
        ending_run._r.insert(0, deepcopy(base_properties))


def _insert_document_review_images(document, image_paths, *, anchor_index):
    """Render reviewed images as reference figures after an introduction paragraph."""
    if not image_paths or len(document.paragraphs) <= anchor_index:
        return

    insertion_point = document.paragraphs[anchor_index]._p
    for position, image_path in enumerate(image_paths, start=1):
        path = Path(image_path)
        if not path.is_file():
            raise OfficialProposalPdfError("Uma imagem da revisao documental nao foi encontrada para gerar o PDF.")

        figure = document.add_table(rows=1, cols=1)
        figure.alignment = WD_TABLE_ALIGNMENT.CENTER
        figure.autofit = False
        cell = figure.cell(0, 0)
        cell.width = Inches(5.45)
        figure.columns[0].width = Inches(5.45)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F7F9FA")
        properties.append(shading)
        margins = OxmlElement("w:tcMar")
        for side in ("top", "left", "bottom", "right"):
            margin = OxmlElement(f"w:{side}")
            margin.set(qn("w:w"), "120")
            margin.set(qn("w:type"), "dxa")
            margins.append(margin)
        properties.append(margins)
        borders = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            border = OxmlElement(f"w:{side}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:color"), "D8E2E5")
            borders.append(border)
        properties.append(borders)

        image_paragraph = cell.paragraphs[0]
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.space_before = Pt(0)
        image_paragraph.paragraph_format.space_after = Pt(4)
        try:
            image_paragraph.add_run().add_picture(str(path), width=Inches(4.95))
        except Exception as error:
            raise OfficialProposalPdfError("Nao foi possivel inserir uma imagem da introducao no documento.") from error
        caption = cell.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(0)
        caption.paragraph_format.space_after = Pt(0)
        caption_run = caption.add_run(f"Figura {position} - Imagem de referencia enviada pelo cliente")
        caption_run.font.name = "Arial"
        caption_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        caption_run.font.size = Pt(8)
        caption_run.font.italic = True
        caption_run.font.color.rgb = RGBColor(92, 108, 123)

        insertion_point.addnext(figure._tbl)
        insertion_point = figure._tbl


def _insert_offshore_intro_images(document, image_paths):
    _insert_document_review_images(document, image_paths, anchor_index=50)


def _insert_pc_onshore_intro_images(document, image_paths):
    _insert_document_review_images(document, image_paths, anchor_index=37)


def _apply_offshore_revision(document, revision, image_paths=None):
    if not revision:
        return
    paragraphs = document.paragraphs
    lines = revision.get("linhas") or {}
    introduction = _clean(revision.get("introducao"))
    procedure_title = _clean(revision.get("procedimentoTitulo"))
    if introduction and len(paragraphs) > 50:
        _set_paragraph_text(paragraphs[50], introduction)
    if procedure_title and len(paragraphs) > 54:
        # The official template uses a static index. Keep it synchronized with
        # the reviewed service title rendered in the procedure section.
        procedure_heading = re.sub(r"^3\.2\s*", "", procedure_title, flags=re.IGNORECASE)
        procedure_heading = f"3.2 {procedure_heading}"
        _set_paragraph_text(paragraphs[24], procedure_heading)
        _set_paragraph_text(paragraphs[54], procedure_heading)
    if len(document.tables) >= 5:
        if lines.get("PROCEDIMENTO"):
            _fill_table_rows(document.tables[2], [(str(index).zfill(2), line.get("descricao", "")) for index, line in enumerate(lines["PROCEDIMENTO"], start=1)], 2)
        if lines.get("EQUIPE"):
            _fill_table_rows(document.tables[3], [(line.get("descricao", ""), line.get("quantidade", "")) for line in lines["EQUIPE"]], 2)
        if lines.get("EQUIPAMENTO"):
            _fill_table_rows(document.tables[4], [(str(index).zfill(2), line.get("descricao", ""), line.get("quantidade", "")) for index, line in enumerate(lines["EQUIPAMENTO"], start=1)], 3)
    for start, content in ((61, lines.get("PREMISSA", [])), (66, lines.get("OBRIGACAO", []))):
        if content:
            for offset in range(max(3, len(content))):
                if start + offset < len(paragraphs):
                    _set_paragraph_text(paragraphs[start + offset], content[offset].get("descricao", "") if offset < len(content) else "")
    # Add visual figures only after the template's indexed paragraphs and
    # tables are complete. A figure is itself a table and would otherwise
    # shift the official table indexes above.
    _insert_offshore_intro_images(document, image_paths)


def _apply_offshore_signature_details(document, analyst, emission_date):
    """Apply the issued-document signature text without retaining the stamp image."""
    paragraphs = document.paragraphs
    if len(paragraphs) <= 82:
        raise OfficialProposalPdfError("A estrutura de assinatura do template Offshore nao foi encontrada.")

    _set_paragraph_text(paragraphs[78], _format_date_long(emission_date or date.today()))
    for drawing in list(paragraphs[79]._p.xpath(".//w:drawing")):
        drawing.getparent().remove(drawing)
    _set_paragraph_text(paragraphs[82], analyst or "Nao informado")

    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            if "Comercial" not in paragraph.text:
                continue
            for run in paragraph.runs:
                if "Comercial" not in run.text:
                    continue
                # Keep the template's leading tab stops: they position the
                # analyst block on the lower-right corner of every page.
                leading_tabs = run.text[:len(run.text) - len(run.text.lstrip("\t"))]
                run.text = f"{leading_tabs}{analyst or 'Nao informado'} - Comercial"
                break


def load_offshore_template_draft():
    """Extract editable defaults from the immutable Offshore template."""
    path = resolve_official_template_path("pc_offshore")
    document = Document(path)
    paragraphs = document.paragraphs
    def rows(table_index, description_col, quantity_col=None):
        result = []
        for row in document.tables[table_index].rows[1:]:
            description = _clean(row.cells[description_col].text)
            if description:
                result.append({"descricao": description, "quantidade": _clean(row.cells[quantity_col].text) if quantity_col is not None else ""})
        return result
    return {
        "introducao": _clean(paragraphs[50].text),
        "procedimentoTitulo": _clean(paragraphs[54].text).replace("3.2 ", "", 1),
        "linhas": {"PROCEDIMENTO": rows(2, 1), "EQUIPE": rows(3, 0, 1), "EQUIPAMENTO": rows(4, 1, 2), "PREMISSA": [{"descricao": _clean(paragraphs[index].text), "quantidade": ""} for index in (61, 62, 63) if _clean(paragraphs[index].text)], "OBRIGACAO": [{"descricao": _clean(paragraphs[index].text), "quantidade": ""} for index in (66, 67) if _clean(paragraphs[index].text)]},
    }


def _assert_no_placeholders(document):
    remaining = "\n".join(paragraph.text for paragraph in _paragraphs(document))
    forbidden = (
        "NOME DA EMPRESA", "NOME DO SOLICITANTE", "E-MAIL DO CLIENTE",
        "Nome cliente", "Nome do solicitante", "e-mail cliente",
        "Nome do serviço", "Nome serviço", "xxxxx", "{{", "}}",
    )
    found = [token for token in forbidden if token.casefold() in remaining.casefold()]
    if found:
        raise OfficialProposalPdfError(
            "O template ainda possui campos pendentes: " + ", ".join(found) + "."
        )


def _convert_with_libreoffice(docx_path, output_dir):
    binary = os.environ.get("LIBREOFFICE_BINARY") or shutil.which("soffice") or shutil.which("libreoffice")
    if not binary:
        return None
    result = subprocess.run(
        [binary, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(docx_path)],
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    output_path = output_dir / f"{docx_path.stem}.pdf"
    if result.returncode != 0 or not output_path.exists():
        raise OfficialProposalPdfError("O LibreOffice não conseguiu converter o documento oficial para PDF.")
    return output_path


def _convert_with_word(docx_path, output_dir):
    """Windows fallback only. Production Linux environments should use LibreOffice."""
    if os.name != "nt":
        return None
    output_path = output_dir / f"{docx_path.stem}.pdf"
    docx_argument = str(docx_path).replace("'", "''")
    pdf_argument = str(output_path).replace("'", "''")
    command = (
        "$ErrorActionPreference='Stop';"
        "$word=New-Object -ComObject Word.Application;"
        "$word.Visible=$false;$word.DisplayAlerts=0;"
        "try {"
        f"$doc=$word.Documents.Open('{docx_argument}', $false, $true);"
        f"$doc.ExportAsFixedFormat('{pdf_argument}',17);"
        "$doc.Close($false)"
        "} finally {$word.Quit()}"
    )
    result = subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-Command", command],
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    if result.returncode != 0 or not output_path.exists():
        raise OfficialProposalPdfError("O Microsoft Word não conseguiu converter o documento oficial para PDF.")
    return output_path


def generate_official_proposal_pdf(
    proposal,
    *,
    serialized,
    items,
    document_revision=None,
    template_key=None,
    preserve_variable_highlights=True,
    document_image_paths=None,
):
    """Fill a private DOCX copy and return the generated PDF bytes and filename."""
    template_key = template_key or _proposal_kind(proposal)
    template_path = resolve_official_template_path(template_key)

    emission_date = proposal.data_emissao
    number = _clean(serialized.get("numeroProposta"))
    revision = _clean(serialized.get("rev")) or "00"
    client = _clean(serialized.get("empresa"))
    service = _clean(serialized.get("servico") or serialized.get("escopo"))
    requester = _clean(serialized.get("solicitante"))
    requester_email = _clean(serialized.get("emailSolicitante"))
    unit = _clean(serialized.get("unidade"))
    responsible = _clean(serialized.get("responsavel"))
    required = {"cliente": client, "serviço": service, "solicitante": requester, "e-mail": requester_email}
    if template_key == "pt_onshore":
        # These cover fields are automatic when the Commercial proposal has
        # them, but the official PT can still be issued without either value.
        required.pop("solicitante", None)
        required.pop("e-mail", None)
    missing = [label for label, value in required.items() if not value]
    if missing:
        raise OfficialProposalPdfError("Preencha os campos obrigatórios do documento: " + ", ".join(missing) + ".")

    replacements = {
        "3140.10.03": number,
        "0027.2025": number,
        "Revisão 00": f"Revisão {revision}",
        "Limpeza de tanque de água potável": service,
        "Nome do serviço": service,
        "Nome serviço": service,
        "nome do serviço": service,
        "FPSO MARIA QUITÉRIA": unit,
        "YINSON": client,
        "NOME DA EMPRESA": client,
        "Nome cliente": client,
        "Ilda Neves": requester,
        "NOME DO SOLICITANTE": requester,
        "Nome do solicitante": requester,
        "ilda.neves@yinson.com": requester_email,
        "E-MAIL DO CLIENTE": requester_email,
        "e-mail cliente": requester_email,
        "Agosto / 2026": _format_cover_month(emission_date),
        "Julho / 2026": _format_cover_month(emission_date),
        "Fernanda Braz – Comercial": f"{responsible} – Comercial",
        "Nome do analista – Comercial": f"{responsible} – Comercial",
        "Katlyn Brito – Comercial": f"{responsible} – Comercial",
        "à nome do cliente,": f"À {client},",
        "à nome cliente,": f"À {client},",
        "BYD_Gerenciamento de Resíduos Classe I - Camaçari_R3": service,
        "Rio de Janeiro, xx de junho de 2026.": _format_date_long(emission_date),
    }

    normalized_items = [
        {
            "id": _clean(item.get("id")),
            "label": _clean(item.get("label") or item.get("nome")),
            "preco_unitario": Decimal(str(item.get("preco_unitario") or "0")),
            "quantidade": Decimal(str(item.get("quantidade") or "1")),
            "unidade": _clean(item.get("unidade")),
        }
        for item in items
    ]
    with tempfile.TemporaryDirectory(prefix=f"synchro_proposta_{number}_") as temporary_dir:
        temporary_dir = Path(temporary_dir)
        copied_template = temporary_dir / f"proposta_{number}_rev_{revision}.docx"
        shutil.copy2(template_path, copied_template)
        document = Document(copied_template)
        # The PT official document explicitly authorizes changes only in yellow
        # runs. Its content must never pass through the broad legacy replacer.
        if template_key != "pt_onshore":
            _replace_document_text(document, replacements)
        if template_key == "pc_offshore":
            _apply_offshore_revision(document, document_revision)
            _fill_offshore_financial_table(document, normalized_items)
            # This figure creates a DOCX table. Insert it only after every
            # indexed template table has been filled.
            _insert_offshore_intro_images(document, document_image_paths)
            _apply_offshore_signature_details(document, responsible, emission_date)
        elif template_key == "pc_onshore":
            if not document_revision:
                raise OfficialProposalPdfError("Não foi possível carregar a revisão da Proposta Comercial Onshore.")
            _apply_pc_onshore_revision(
                document,
                document_revision,
                serialized,
                normalized_items,
                show_variable_highlights=preserve_variable_highlights,
            )
            _insert_pc_onshore_intro_images(document, document_image_paths)
            _append_pc_onshore_signature(document, responsible)
            _set_pc_onshore_footer_analyst(document, responsible)
        elif template_key == "pt_onshore":
            if not document_revision:
                raise OfficialProposalPdfError("Nao foi possivel carregar a revisao da Proposta Tecnica Onshore.")
            _apply_pt_onshore_revision(
                document,
                document_revision,
                serialized,
                show_variable_highlights=preserve_variable_highlights,
            )
        if not preserve_variable_highlights and template_key != "pt_onshore":
            _remove_text_highlights(document)
        if template_key != "pt_onshore":
            _assert_no_placeholders(document)
        document.save(copied_template)

        output_path = _convert_with_libreoffice(copied_template, temporary_dir)
        if output_path is None:
            output_path = _convert_with_word(copied_template, temporary_dir)
        if output_path is None:
            raise OfficialProposalPdfError(
                "Nenhum conversor DOCX para PDF está disponível. Configure LibreOffice em modo headless."
            )
        return output_path.read_bytes(), f"{number}_{template_key.upper()}_REV{revision}.pdf"
