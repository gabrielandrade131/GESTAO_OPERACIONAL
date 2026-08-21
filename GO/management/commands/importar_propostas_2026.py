"""Importa propostas históricas de 2026 de forma auditável e idempotente."""

import csv
import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

from GO.models import (
    Cliente,
    Financeiro,
    MetodoOperacional,
    OrdemServico,
    RdoTanque,
    ResponsavelCoordenador,
    ServicoComercial,
    Unidade,
)


SHEET_NAME = "2026"
REPORT_FILENAME = "importacao_propostas_2026_relatorio.csv"

HEADERS = {
    "numero": "Nº de Proposta",
    "revisao": "REV",
    "emissao": "Emissão",
    "responsavel": "Quem enviou?",
    "entrega": "Data de entrega da proposta",
    "solicitacao": "Data de solicitação da proposta",
    "fechamento": "Data de Fechamento",
    "previsao": "Previsão de contratação",
    "follow_up": "FOLLOW UP",
    "natureza": "Natureza",
    "tipo_operacao": "Unidade",
    "heat_map": "HEAT MAP",
    "status": "Status Proposta",
    "motivo": "Motivo (Declínio ou Perda)",
    "pt": "PT",
    "pc": "PC/PTC",
    "cliente": "Empresa",
    "uf": "UF",
    "local": "EMBARCAÇÃO/LOCAL",
    "servico": "Escopo",
    "receita": "Estimativa Receita",
    "tempo": "Tempo de contrato (em dias)",
    "solicitante": "Solicitante",
    "fonte": "Fonte do Lead",
    "comentario": "Comentário",
    "segmento": "Segmento Cliente",
}


def normalize(value):
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return text.casefold()


def clean(value):
    # REV 00 and zero-valued monetary/quantity cells are valid spreadsheet values.
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_revision(value):
    match = re.search(r"\d+", clean(value))
    return int(match.group()) if match else None


def parse_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        # Excel serial dates are only accepted in a conservative range.
        if 20000 <= value <= 70000:
            return date(1899, 12, 30).fromordinal(date(1899, 12, 30).toordinal() + int(value))
        return None
    text = clean(value)
    for pattern in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def parse_decimal(value):
    if value in (None, ""):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = clean(value).replace("R$", "").replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def parse_integer(value):
    match = re.search(r"\d+", clean(value))
    return int(match.group()) if match else None


def normalize_choice(value, choices, aliases=None):
    raw = clean(value)
    if not raw:
        return ""
    aliases = aliases or {}
    lookup = {normalize(choice): choice for choice in choices}
    lookup.update({normalize(key): target for key, target in aliases.items()})
    return lookup.get(normalize(raw), "")


class Command(BaseCommand):
    help = "Importa propostas históricas da aba 2026, com dry-run obrigatório por padrão."

    def add_arguments(self, parser):
        parser.add_argument("--dry-run", action="store_true", help="Processa e gera relatório sem gravar dados.")
        parser.add_argument("--execute", action="store_true", help="Persiste os registros validados.")
        parser.add_argument("--arquivo", type=str, help="Caminho relativo ao projeto para a planilha.")
        parser.add_argument("--relatorio", type=str, help="Caminho relativo para o CSV de auditoria.")

    def handle(self, *args, **options):
        if options["dry_run"] and options["execute"]:
            raise CommandError("Use somente um entre --dry-run e --execute.")
        execute = bool(options["execute"])
        file_path = self._resolve_file(options.get("arquivo"))
        report_path = self._resolve_report(options.get("relatorio"))
        rows = self._read_rows(file_path)
        selected, reports, summary = self._select_winners(rows)
        summary["rows_read"] = len(rows)
        summary["candidates_after_revision"] = len(selected)
        self._validate_candidates(selected, reports, summary)
        self._write_report(report_path, reports)

        if not execute:
            self._print_summary(file_path, report_path, summary, dry_run=True)
            return

        if summary["blocking_errors"]:
            self.stdout.write(
                self.style.WARNING(
                    "Há linhas bloqueadas; elas serão ignoradas e as propostas prontas serão importadas."
                )
            )

        self._persist(selected, reports, summary)
        self._write_report(report_path, reports)
        self._print_summary(file_path, report_path, summary, dry_run=False)
        self.stdout.write(self.style.SUCCESS(f"Importação concluída. Relatório: {report_path}"))

    def _resolve_file(self, raw_path):
        if raw_path:
            candidate = Path(raw_path)
            if not candidate.is_absolute():
                candidate = Path(settings.BASE_DIR) / candidate
            if candidate.exists():
                return candidate
            raise CommandError(f"Arquivo não encontrado: {candidate}")

        directories = [
            Path(settings.BASE_DIR) / "GO" / "docs" / "planilha_propostas",
            Path(settings.BASE_DIR) / "GO" / "docs",
        ]
        matches = []
        for directory in directories:
            if directory.exists():
                matches.extend(directory.glob("*Controle*Propostas*.xlsx"))
        if not matches:
            raise CommandError("Planilha de propostas não encontrada em GO/docs/planilha_propostas ou GO/docs.")
        return sorted(matches)[0]

    def _resolve_report(self, raw_path):
        if raw_path:
            return Path(settings.BASE_DIR) / raw_path
        return Path(settings.BASE_DIR) / "GO" / "docs" / "planilha_propostas" / REPORT_FILENAME

    def _read_rows(self, file_path):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError("openpyxl é obrigatório. Instale as dependências do projeto.") from exc

        workbook = load_workbook(file_path, read_only=True, data_only=True)
        if SHEET_NAME not in workbook.sheetnames:
            raise CommandError(f"A aba obrigatória '{SHEET_NAME}' não existe na planilha.")
        worksheet = workbook[SHEET_NAME]
        header_row, columns = self._find_headers(worksheet)
        missing = [label for key, label in HEADERS.items() if key not in columns and key not in {"comentario"}]
        if missing:
            raise CommandError(f"Colunas obrigatórias não encontradas: {', '.join(missing)}")

        records = []
        for excel_row, values in enumerate(worksheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            raw = {key: values[index] if index < len(values) else None for key, index in columns.items()}
            if not any(value not in (None, "") for value in raw.values()):
                continue
            number = parse_integer(raw.get("numero"))
            revision = parse_revision(raw.get("revisao"))
            if number is None or revision is None:
                invalid_parts = []
                if number is None:
                    invalid_parts.append("Número da proposta ausente ou inválido")
                if revision is None:
                    invalid_parts.append("REV ausente ou inválida")
                records.append({"excel_row": excel_row, "raw": raw, "invalid_reason": "; ".join(invalid_parts) + "."})
                continue
            records.append({
                "excel_row": excel_row,
                "raw": raw,
                "number": number,
                "revision": revision,
                "emission": parse_date(raw.get("emissao")),
                "delivery": parse_date(raw.get("entrega")),
                "closing": parse_date(raw.get("fechamento")),
                "forecast": parse_date(raw.get("previsao")),
                "request": parse_date(raw.get("solicitacao")),
                "identity": (number, normalize(raw.get("cliente")), normalize(raw.get("local")), normalize(raw.get("tipo_operacao"))),
            })
        workbook.close()
        return records

    def _find_headers(self, worksheet):
        for row_number, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
            normalized = {normalize(value): index for index, value in enumerate(values) if clean(value)}
            if normalize(HEADERS["numero"]) not in normalized or normalize(HEADERS["revisao"]) not in normalized:
                continue
            columns = {}
            for key, label in HEADERS.items():
                index = normalized.get(normalize(label))
                if index is not None:
                    columns[key] = index
            return row_number, columns
        raise CommandError("Cabeçalho da planilha não encontrado na aba 2026.")

    def _selection_key(self, record):
        fallback_dates = [record.get(key) for key in ("delivery", "closing", "forecast", "request")]
        fallback = max((item for item in fallback_dates if item), default=date.min)
        return (record["revision"], record.get("emission") or date.min, fallback, record["excel_row"])

    def _select_winners(self, rows):
        reports, summary, groups = [], Counter(), defaultdict(list)
        for record in rows:
            if record.get("invalid_reason"):
                reports.append(self._report(record, "BLOCKER", record["invalid_reason"]))
                summary["blocking_errors"] += 1
                continue
            if record["emission"] is None:
                reports.append(self._report(record, "IGNORADO", "Emissão vazia ou inválida."))
                summary["invalid_dates"] += 1
                continue
            if record["emission"].year != 2026:
                reports.append(self._report(record, "IGNORADO", "Emissão fora de 2026."))
                summary["outside_year"] += 1
                continue
            groups[record["identity"]].append(record)

        selected = []
        for identity, group in groups.items():
            winner = max(group, key=self._selection_key)
            selected.append(winner)
            summary["valid_rows"] += len(group)
            if len(group) > 1:
                summary["duplicate_groups"] += 1
                summary["discarded_revisions"] += len(group) - 1
                for record in group:
                    if record is not winner:
                        reports.append(self._report(record, "REV DESCARTADA", f"Venceu a linha {winner['excel_row']} (REV {winner['revision']:02d}).", winner))

        same_number = defaultdict(list)
        for record in selected:
            same_number[record["number"]].append(record)
        for number, group in same_number.items():
            if len(group) > 1:
                summary["reused_numbers"] += 1
                for record in group:
                    record.setdefault("warnings", []).append(
                        f"número {number} reutilizado por oportunidades com identidades diferentes"
                    )
        return selected, reports, summary

    def _validate_candidates(self, selected, reports, summary):
        status_values = [value for value, _label in Financeiro._meta.get_field("status_proposta").choices]
        natureza_values = [value for value, _label in Financeiro._meta.get_field("natureza").choices]
        fonte_values = [value for value, _label in Financeiro._meta.get_field("fonte_lead").choices]
        segmento_values = [value for value, _label in Financeiro._meta.get_field("segmento_cliente").choices]
        motivo_values = [value for value, _label in Financeiro._meta.get_field("motivo_perda").choices]
        pt_values = [value for value, _label in Financeiro._meta.get_field("pt_financeiro").choices]
        pc_values = [value for value, _label in Financeiro._meta.get_field("pc_ptc").choices]
        uf_values = [value for value, _label in Financeiro._meta.get_field("uf").choices]

        aliases = {
            "natureza": {"contrato novo": "Contrato Novo", "contrato novo ": "Contrato Novo"},
            "fonte": {"convite direto": "Convite Direto", "cross sell": "Cross Shell", "vendas ambipar": "Vendas Ambipar"},
            "motivo": {"enviado outra unid ambipar": "Enviado outra unidade AMBIPAR"},
        }
        state_aliases = {"rio de janeiro": "RJ", "sao paulo": "SP", "minas gerais": "MG", "mato grosso": "MT", "santa catarina": "SC", "pernambuco": "PE", "bahia": "BA", "espirito santo": "ES", "parana": "PR", "goias": "GO", "ceara": "CE", "para": "PA", "amazonas": "AM", "rio grande do sul": "RS"}

        for record in selected:
            raw = record["raw"]
            record["natureza"] = normalize_choice(raw.get("natureza"), natureza_values, aliases["natureza"])
            record["status"] = normalize_choice(raw.get("status"), status_values)
            record["fonte"] = normalize_choice(raw.get("fonte"), fonte_values, aliases["fonte"])
            record["segmento"] = normalize_choice(raw.get("segmento"), segmento_values)
            record["motivo"] = normalize_choice(raw.get("motivo"), motivo_values, aliases["motivo"])
            record["pt"] = normalize_choice(raw.get("pt"), pt_values)
            record["pc"] = normalize_choice(raw.get("pc"), pc_values)
            record["uf"] = normalize_choice(raw.get("uf"), uf_values, state_aliases)
            record["tipo_operacao"] = normalize_choice(raw.get("tipo_operacao"), ["Onshore", "Offshore"])
            record["receita"] = parse_decimal(raw.get("receita"))
            record["tempo"] = parse_integer(raw.get("tempo"))
            record["errors"] = []
            required = {
                "tipo de operação": record["tipo_operacao"],
                "responsável": clean(raw.get("responsavel")),
            }
            for label, value in required.items():
                if not value:
                    record["errors"].append(f"{label} ausente ou não reconhecido")
            record.setdefault("warnings", [])
            optional_values = {
                "estimativa de receita": record["receita"],
                "UF": record["uf"],
                "segmento": record["segmento"],
                "fonte do lead": record["fonte"],
                "natureza": record["natureza"],
                "status": record["status"],
                "cliente": clean(raw.get("cliente")),
                "embarcacao/local": clean(raw.get("local")),
            }
            for label, value in optional_values.items():
                if not value:
                    record["warnings"].append(f"{label} ausente ou não reconhecido no histórico")
            service_name = clean(raw.get("servico"))
            if service_name and len(service_name) > ServicoComercial._meta.get_field("nome").max_length:
                record["warnings"].append("serviço excede o limite do catálogo e será preservado somente na proposta histórica")
            if not record["motivo"] and clean(raw.get("motivo")):
                record["warnings"].append("motivo de perda não reconhecido no histórico")
            summary["warnings"] += len(record["warnings"])
            if record["errors"]:
                summary["blocking_errors"] += 1
                reports.append(self._report(record, "BLOCKER", "; ".join(record["errors"])))
            elif self._historical_record_exists(record):
                record["existing"] = True
                summary["existing"] += 1
                reports.append(self._report(record, "JÁ EXISTENTE", "A mesma oportunidade histórica já existe no Synchro."))
            else:
                summary["ready"] += 1
                if record["warnings"]:
                    reports.append(self._report(record, "PRONTA", "Validada para importação.", warnings=record["warnings"]))
                else:
                    reports.append(self._report(record, "PRONTA", "Validada para importação."))

    def _persist(self, selected, reports, summary):
        base_os = OrdemServico.objects.order_by("id").first()
        tank = RdoTanque.objects.order_by("id").first()
        if base_os is None or tank is None:
            raise CommandError("A importação requer pelo menos uma OS e um tanque de RDO como referências técnicas legadas.")

        with transaction.atomic():
            for record in selected:
                if record.get("errors") or record.get("existing"):
                    continue
                raw = record["raw"]
                cliente_nome = clean(raw.get("cliente"))
                local_nome = clean(raw.get("local"))
                cliente = self._get_or_create(Cliente, cliente_nome, summary, "clientes_created") if cliente_nome else None
                unidade = self._get_or_create(Unidade, local_nome, summary, "units_created") if local_nome else None
                responsavel = self._get_or_create_responsavel(clean(raw.get("responsavel")), summary)
                servico = clean(raw.get("servico"))
                if servico and len(servico) <= ServicoComercial._meta.get_field("nome").max_length:
                    self._get_or_create(ServicoComercial, servico, summary, "services_created", ativo=True)
                method, _ = MetodoOperacional.objects.get_or_create(nome="N/A", defaults={"ativo": True})
                request_date = record.get("request") or record["emission"]
                forecast_date = record.get("forecast") or record.get("delivery") or request_date
                followup_date = parse_date(raw.get("follow_up"))
                bundle = {
                    "overrides": {
                        "empresa": cliente_nome,
                        "unidade": local_nome,
                        "embarcacao_local": local_nome,
                        "tipo_operacao": record["tipo_operacao"],
                        "servico": servico,
                    },
                    "items": [],
                }
                if followup_date:
                    bundle["items"].append({"data": followup_date.strftime("%d/%m/%Y"), "dataProximaAcao": followup_date.strftime("%d/%m/%Y"), "responsavel": responsavel.nome, "status": "Pendente", "comentario": "Follow-up histórico importado.", "proximaAcao": "Follow-up histórico importado."})

                Financeiro.objects.create(
                    proposta=record["number"], revisao=record["revision"], data_emissao=record["emission"],
                    data_solicitacao_proposta=request_date, data_fechamento_proposta=record.get("closing"),
                    previsao_contratacao=forecast_date, follow_up=__import__("json").dumps(bundle, ensure_ascii=False),
                    natureza=record["natureza"], heat_map=parse_integer(raw.get("heat_map")) or 0,
                    motivo_perda=record["motivo"] or None, po="", rfi="", cliente=base_os if cliente else None, unidade=base_os,
                    solicitante=clean(raw.get("solicitante")), tipo_operacao=base_os, metodo=base_os,
                    metodo_cadastro=method, data_inicio_frente=base_os, data_fim=base_os, data_fim_frente=base_os,
                    data_entrega_proposta=record.get("delivery"), tempo_contrato_dias=record["tempo"],
                    status_proposta=record["status"], cordenador=base_os, responsavel=responsavel.nome,
                    responsavel_cadastro=responsavel,
                    servico=servico if len(servico) <= Financeiro._meta.get_field("servico").max_length else None,
                    volume_tanque_exec=tank,
                    comentario=clean(raw.get("comentario")), requisitos_cliente="", requisitos_ambipar="", treinamentos="", ajuste_operacional="",
                    analise_critica=False, pt_financeiro=record["pt"] or "Pendente", pc_ptc=record["pc"] or "Pendente",
                    uf=record["uf"] or None, estimativo_receita=record["receita"], fonte_lead=record["fonte"] or None, segmento_cliente=record["segmento"] or None,
                    importado_historico=True,
                )
                record["imported"] = True
                summary["imported"] += 1
                self._mark_imported_report(reports, record)

    def _get_or_create(self, model, name, summary, counter_key, **defaults):
        instance = model.objects.filter(nome__iexact=name).first()
        if instance:
            return instance
        instance = model.objects.create(nome=name, **defaults)
        summary[counter_key] += 1
        return instance

    def _historical_record_exists(self, record):
        """Compare historical rows by their logical opportunity identity, not by number only."""
        expected = record["identity"]
        for proposal in Financeiro.objects.filter(proposta=record["number"]).only("proposta", "follow_up"):
            try:
                bundle = json.loads(proposal.follow_up or "{}")
            except (TypeError, ValueError):
                bundle = {}
            overrides = bundle.get("overrides") if isinstance(bundle, dict) else {}
            overrides = overrides if isinstance(overrides, dict) else {}
            actual = (
                proposal.proposta,
                normalize(overrides.get("empresa")),
                normalize(overrides.get("embarcacao_local") or overrides.get("unidade")),
                normalize(overrides.get("tipo_operacao")),
            )
            if actual == expected:
                return True
        return False

    def _get_or_create_responsavel(self, name, summary):
        person = ResponsavelCoordenador.objects.filter(nome__iexact=name).first()
        if person:
            if not person.responsavel_comercial:
                person.responsavel_comercial = True
                person.save(update_fields=["responsavel_comercial", "atualizado_em"])
            return person
        summary["responsaveis_created"] += 1
        return ResponsavelCoordenador.objects.create(nome=name, responsavel_comercial=True, ativo=True)

    def _report(self, record, status, reason, winner=None, warnings=None):
        raw = record.get("raw", {})
        return {
            "linha_excel": record.get("excel_row", ""), "numero_proposta": record.get("number", raw.get("numero", "")),
            "revisao": record.get("revision", raw.get("revisao", "")), "cliente": clean(raw.get("cliente")),
            "status_importacao": status, "motivo": reason,
            "linha_selecionada": winner.get("excel_row", "") if winner else record.get("excel_row", ""),
            "severidade": "WARNING" if warnings else ("BLOCKER" if status == "BLOCKER" else "OK"),
            "observacao": "; ".join(warnings or []),
        }

    def _mark_imported_report(self, reports, record):
        for report in reports:
            if report["linha_excel"] == record["excel_row"] and report["status_importacao"] == "PRONTA":
                report["status_importacao"] = "IMPORTADA"
                report["motivo"] = "Importada sem sobrescrever registros existentes."
                return

    def _write_report(self, report_path, reports):
        report_path.parent.mkdir(parents=True, exist_ok=True)
        fields = ["linha_excel", "numero_proposta", "revisao", "cliente", "status_importacao", "severidade", "motivo", "linha_selecionada", "observacao"]
        with report_path.open("w", newline="", encoding="utf-8-sig") as file_handle:
            writer = csv.DictWriter(file_handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows(reports)

    def _print_summary(self, file_path, report_path, summary, dry_run):
        self.stdout.write("=" * 46)
        self.stdout.write("IMPORTACAO 2026")
        self.stdout.write("=" * 46)
        self.stdout.write(f"Linhas analisadas: {summary['rows_read']}")
        self.stdout.write(f"Linhas de 2026: {summary['valid_rows']}")
        self.stdout.write(f"Propostas apos deduplicacao: {summary['candidates_after_revision']}")
        self.stdout.write(f"Warnings: {summary['warnings']}")
        self.stdout.write(f"Blockers: {summary['blocking_errors']}")
        self.stdout.write(f"Prontas para importar: {summary['ready']}")
        self.stdout.write(f"Ignoradas por blocker: {summary['blocking_errors']}")
        self.stdout.write(f"Revisoes antigas descartadas: {summary['discarded_revisions']}")
        self.stdout.write(f"Fora de 2026: {summary['outside_year']}")
        self.stdout.write(f"Numeros historicos reutilizados: {summary['reused_numbers']}")
        self.stdout.write(f"Ja existentes: {summary['existing']}")
        self.stdout.write(f"Importadas: {summary['imported']}")
        self.stdout.write(f"Relatorio: {report_path}")
